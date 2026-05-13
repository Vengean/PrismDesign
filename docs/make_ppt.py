#!/usr/bin/env python3
"""Generate PrismDesign PPT v2 — 3 pages, white theme."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Brand colors ──
PRIMARY = RGBColor(99, 102, 241)
DARK = RGBColor(30, 30, 46)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(107, 114, 128)
LIGHT_GRAY = RGBColor(240, 240, 248)
BLUE = RGBColor(59, 130, 246)
PURPLE = RGBColor(139, 92, 246)
GREEN = RGBColor(34, 197, 94)
RED = RGBColor(239, 68, 68)
ORANGE = RGBColor(245, 158, 11)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text_box(slide, left, top, width, height, text, size=18,
             color=DARK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Microsoft YaHei"
    p.alignment = align
    return tb


def rect(slide, left, top, width, height, fill, text="",
         size=14, fc=WHITE, bold=False, align=PP_ALIGN.CENTER):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if text:
        tf = s.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = fc
        p.font.bold = bold
        p.font.name = "Microsoft YaHei"
        p.alignment = align
    return s


def bullets(slide, left, top, width, items, size=13, color=DARK, spacing=Pt(6)):
    tb = slide.shapes.add_textbox(left, top, width, Inches(len(items) * 0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Microsoft YaHei"
        p.space_after = spacing


# ============================================================
# Page 1: 当前 AI 协作模式的痛点
# ============================================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s1, WHITE)

text_box(s1, Inches(1.0), Inches(0.5), Inches(11), Inches(0.8),
         "01  当前 AI 协作模式的痛点", 36, DARK, bold=True)
text_box(s1, Inches(1.0), Inches(1.25), Inches(11), Inches(0.5),
         "AI 时代，研发流程的效率瓶颈不在编码，而在协作", 18, GRAY)

# ── Flow diagram: 3 role cards in a horizontal strip ──
flow_y = Inches(2.1)
strip_x = Inches(1.0)
strip_total = Inches(11.3)
step_w = Inches(3.3)
arrow_w = Inches(0.7)

flow_roles = [
    ("产品经理", "用 AI 生成 HTML 原型", BLUE),
    ("设计师", "用 AI 生成独立设计稿", PURPLE),
    ("前端开发", "用 AI 从零写项目代码", GREEN),
]
for i, (role, action, color) in enumerate(flow_roles):
    sx = strip_x + i * (step_w + arrow_w)
    # Step card
    rect(s1, sx, flow_y, step_w, Inches(0.75), RGBColor(248, 248, 252))
    # Color left accent bar
    rect(s1, sx, flow_y, Inches(0.06), Inches(0.75), color)
    # Role name
    text_box(s1, sx + Inches(0.2), flow_y + Inches(0.05), step_w - Inches(0.3), Inches(0.35),
             role, 14, color, bold=True)
    # Action
    text_box(s1, sx + Inches(0.2), flow_y + Inches(0.38), step_w - Inches(0.3), Inches(0.3),
             action, 12, GRAY)
    # Arrow between steps
    if i < 2:
        ax = sx + step_w + Inches(0.15)
        text_box(s1, ax, flow_y + Inches(0.15), Inches(0.4), Inches(0.45),
                 "→", 22, RGBColor(200, 200, 210), bold=True, align=PP_ALIGN.CENTER)

# ── Pain points — 2x2 grid ──
card_y_row1 = Inches(3.2)
card_w = Inches(5.5)
card_h = Inches(1.5)
card_gap_x = Inches(0.35)
card_gap_y = Inches(0.25)
card_start_x = (prs.slide_width - card_w * 2 - card_gap_x) / 2

pains = [
    ("上下文断裂", BLUE, [
        "各角色独立生成，产出物无法复用",
        "无法继承项目已有的组件库和设计规范",
        "AI产出不符合项目技术栈",
    ]),
    ("沟通效率低", PURPLE, [
        "与AI：用纯文字描述需求非常低效且容易出错",
        "与人：需求、设计、开发多轮会议反复对齐",
    ]),
    ("重复劳动", ORANGE, [
        "产品维护独立原型，与代码不同步",
        "设计维护独立设计稿，走查反复修改",
        "同一个 UI 被做了三遍，一旦有调整三处都要改",
    ]),
    ("环境割裂", RED, [
        "缺乏统一的MCP和skill管理",
        "设计规范和组件库在设计软件中难以被AI读取",
    ]),
]

for i, (title, color, items) in enumerate(pains):
    col = i % 2
    row = i // 2
    cx = card_start_x + col * (card_w + card_gap_x)
    cy = card_y_row1 + row * (card_h + card_gap_y)

    # Card background
    rect(s1, cx, cy, card_w, card_h, RGBColor(248, 248, 252))
    # Left accent bar
    rect(s1, cx, cy, Inches(0.05), card_h, color)

    # Number badge + title
    rect(s1, cx + Inches(0.2), cy + Inches(0.2), Inches(0.32), Inches(0.32),
         color, str(i + 1), 12, WHITE, bold=True)
    text_box(s1, cx + Inches(0.62), cy + Inches(0.2), Inches(2.0), Inches(0.35),
             title, 15, DARK, bold=True)

    # Items
    for j, item in enumerate(items):
        iy = cy + Inches(0.65) + j * Inches(0.28)
        rect(s1, cx + Inches(0.3), iy + Inches(0.08), Inches(0.05), Inches(0.05), color)
        text_box(s1, cx + Inches(0.5), iy, card_w - Inches(0.8), Inches(0.28),
                 item, 11, RGBColor(80, 85, 105))

# ── Bottom summary ──
rect(s1, Inches(2.2), Inches(6.55), Inches(8.9), Inches(0.55),
     RGBColor(240, 238, 255),
     "核心问题：各角色独立作业，产出物无法衔接，AI 能力被碎片化使用",
     14, PRIMARY, bold=True)


# ============================================================
# Page 2: 棱镜给每个岗位带来的效率提升
# ============================================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s2, WHITE)

text_box(s2, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
         "02  棱镜带来的效率提升", 36, DARK, bold=True)
text_box(s2, Inches(0.8), Inches(1.2), Inches(11), Inches(0.5),
         "每个角色都在自己的专业领域内工作得更好，同时所有工作汇聚到同一份代码", 18, GRAY)

# Three role columns
col_w = Inches(3.7)
col_gap = Inches(0.3)
col_start = Inches(0.8)
col_y = Inches(2.0)

role_data = [
    {
        "title": "产品经理",
        "color": BLUE,
        "solved": "解决的痛点",
        "items": [
            ("选中元素精准沟通", "点击即定位，告别「那个按钮」的模糊描述"),
            ("评论标注需求", "对着具体元素写修改意图，Agent 零歧义执行"),
            ("复用mock/测试/生产数据", "不用再关心造数问题"),
            ("不再维护独立原型", "源码运行起来就是最新产品形态，始终同步"),
        ]
    },
    {
        "title": "设计师",
        "color": PURPLE,
        "solved": "解决的痛点",
        "items": [
            ("可视化编辑样式", "直接调颜色、字号、间距，实时预览效果"),
            ("基于真实组件库工作", "设计规范和 Token 天然一致，不会偏离"),
            ("视觉走查消失", "所见即代码，不存在「还原不一致」"),
            ("不再维护独立设计稿", "在真实页面上调整，告别 Figma 同步维护"),
        ]
    },
    {
        "title": "前端开发",
        "color": GREEN,
        "solved": "解决的痛点",
        "items": [
            ("拿到规范代码骨架", "Agent 生成的代码符合项目技术栈和组件规范"),
            ("只需补业务逻辑", "不再从零编写页面 UI 和样式还原"),
            ("统一维护 Agent 环境", "集中配置 skill / MCP / 模型，全团队共享"),
            ("Git 天然版本管理", "所有修改可追溯、可回滚、可 Code Review"),
        ]
    },
]

for i, role in enumerate(role_data):
    cx = col_start + i * (col_w + col_gap)

    # Column background
    rect(s2, cx, col_y, col_w, Inches(4.8), RGBColor(248, 248, 252))
    # Role title
    rect(s2, cx, col_y, col_w, Inches(0.5), role["color"], role["title"], 16, WHITE, bold=True)

    # Items
    for j, (title, desc) in enumerate(role["items"]):
        iy = col_y + Inches(0.7) + j * Inches(1.0)
        # Item title
        text_box(s2, cx + Inches(0.25), iy, col_w - Inches(0.5), Inches(0.35),
                 title, 13, role["color"], bold=True)
        # Item description
        text_box(s2, cx + Inches(0.25), iy + Inches(0.3), col_w - Inches(0.5), Inches(0.5),
                 desc, 11, RGBColor(100, 100, 115))

# Bottom
rect(s2, Inches(1.5), Inches(6.95), Inches(10.3), Inches(0.45),
     RGBColor(240, 238, 255),
     "不是让所有人写代码，而是让每个角色用最擅长的方式工作，最终汇聚到同一份代码上",
     13, PRIMARY, bold=True)


# ============================================================
# Page 3: 实现方案
# ============================================================
s3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s3, WHITE)

text_box(s3, Inches(1.0), Inches(0.5), Inches(11), Inches(0.8),
         "03  实现方案", 36, DARK, bold=True)
text_box(s3, Inches(1.0), Inches(1.2), Inches(11), Inches(0.5),
         "Chrome 插件 + AI Agent 服务，两个组件协同工作", 18, GRAY)

# ── Architecture diagram — single horizontal row, evenly spaced ──
# [User] → [Chrome Plugin + Running Page] ↔ [Agent Service] → [Project Code]

dia_y = Inches(2.2)
box_w = Inches(2.6)      # uniform width for all 4 boxes
box_h = Inches(2.4)      # uniform height for feature-list boxes
arrow_w = Inches(0.6)    # space for arrows
total_w = box_w * 4 + arrow_w * 3
start_x = (prs.slide_width - total_w) / 2

def dia_box_x(i):
    return start_x + i * (box_w + arrow_w)

def dia_arrow(s, i, label="", bidir=False):
    """Draw arrow between box i and box i+1, centered in the gap."""
    gap_left = dia_box_x(i) + box_w
    gap_right = dia_box_x(i + 1)
    gap_center_x = (gap_left + gap_right) / 2
    mid_y = dia_y + box_h / 2
    shaft_thick = Inches(0.025)
    head_size = Inches(0.12)
    margin = Inches(0.08)  # gap from box edge

    if bidir:
        shaft_left = gap_left + margin + head_size
        shaft_right = gap_right - margin - head_size
        # Left arrowhead
        tl = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
             gap_left + margin, mid_y - head_size / 2, head_size, head_size)
        tl.fill.solid(); tl.fill.fore_color.rgb = RGBColor(180, 185, 200)
        tl.line.fill.background(); tl.rotation = 270.0
        # Shaft
        shaft = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
             shaft_left, mid_y - shaft_thick / 2, shaft_right - shaft_left, shaft_thick)
        shaft.fill.solid(); shaft.fill.fore_color.rgb = RGBColor(180, 185, 200)
        shaft.line.fill.background()
        # Right arrowhead
        tr = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
             gap_right - margin - head_size, mid_y - head_size / 2, head_size, head_size)
        tr.fill.solid(); tr.fill.fore_color.rgb = RGBColor(180, 185, 200)
        tr.line.fill.background(); tr.rotation = 90.0
    else:
        shaft_left = gap_left + margin
        shaft_right = gap_right - margin - head_size
        # Shaft
        shaft = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
             shaft_left, mid_y - shaft_thick / 2, shaft_right - shaft_left, shaft_thick)
        shaft.fill.solid(); shaft.fill.fore_color.rgb = RGBColor(180, 185, 200)
        shaft.line.fill.background()
        # Right arrowhead
        tr = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
             shaft_right, mid_y - head_size / 2, head_size, head_size)
        tr.fill.solid(); tr.fill.fore_color.rgb = RGBColor(180, 185, 200)
        tr.line.fill.background(); tr.rotation = 90.0

    if label:
        text_box(s, gap_left, mid_y + Inches(0.15), gap_right - gap_left, Inches(0.25),
                 label, 9, GRAY, align=PP_ALIGN.CENTER)

# Box 1: User
bx = dia_box_x(0)
rect(s3, bx, dia_y, box_w, box_h, RGBColor(248, 248, 252))
rect(s3, bx, dia_y, box_w, Inches(0.4), GRAY, "用户", 13, WHITE, bold=True)
user_items = ["产品经理", "设计师", "前端开发"]
for j, item in enumerate(user_items):
    text_box(s3, bx + Inches(0.3), dia_y + Inches(0.55) + j * Inches(0.35),
             box_w - Inches(0.6), Inches(0.3), item, 12, RGBColor(80, 85, 105), align=PP_ALIGN.CENTER)

dia_arrow(s3, 0, "浏览器操作")

# Box 2: Chrome Plugin
bx = dia_box_x(1)
rect(s3, bx, dia_y, box_w, box_h, RGBColor(248, 248, 252))
rect(s3, bx, dia_y, box_w, Inches(0.4), PRIMARY, "Chrome 插件", 13, WHITE, bold=True)
plug_items = ["浮动工具栏", "元素选中与检测", "属性编辑面板", "评论标注", "待同步面板"]
for j, item in enumerate(plug_items):
    text_box(s3, bx + Inches(0.2), dia_y + Inches(0.5) + j * Inches(0.28),
             box_w - Inches(0.4), Inches(0.25), f"· {item}", 10, RGBColor(80, 85, 105))

dia_arrow(s3, 1, "WebSocket", bidir=True)

# Box 3: Agent Service
bx = dia_box_x(2)
rect(s3, bx, dia_y, box_w, box_h, RGBColor(248, 248, 252))
rect(s3, bx, dia_y, box_w, Inches(0.4), DARK, "AI Agent 服务", 13, WHITE, bold=True)
agt_items = ["项目扫描", "上下文构建", "代码搜索与读写", "HMR 热更新联动", "MCP / Skill 扩展"]
for j, item in enumerate(agt_items):
    text_box(s3, bx + Inches(0.2), dia_y + Inches(0.5) + j * Inches(0.28),
             box_w - Inches(0.4), Inches(0.25), f"· {item}", 10, RGBColor(80, 85, 105))

dia_arrow(s3, 2, "读写代码")

# Box 4: Project Code
bx = dia_box_x(3)
rect(s3, bx, dia_y, box_w, box_h, RGBColor(248, 248, 252))
rect(s3, bx, dia_y, box_w, Inches(0.4), GREEN, "项目源代码", 13, WHITE, bold=True)
code_items = ["组件库", "设计规范 / Token", "业务代码", "Git 版本管理"]
for j, item in enumerate(code_items):
    text_box(s3, bx + Inches(0.2), dia_y + Inches(0.5) + j * Inches(0.28),
             box_w - Inches(0.4), Inches(0.25), f"· {item}", 10, RGBColor(80, 85, 105))

# ── "Running Page" below — spans from Plugin to Code ──
plug_bx = dia_box_x(1)
code_bx = dia_box_x(3)
page_x = plug_bx
page_w = code_bx + box_w - plug_bx  # span from Plugin left to Code right
page_h = Inches(1.1)
page_y = dia_y + box_h + Inches(0.6)

rect(s3, page_x, page_y, page_w, page_h, RGBColor(248, 248, 252))
rect(s3, page_x, page_y, page_w, Inches(0.4), BLUE, "运行中的前端页面", 13, WHITE, bold=True)
text_box(s3, page_x + Inches(0.2), page_y + Inches(0.5), page_w - Inches(0.4), Inches(0.5),
         "插件通过 Content Script 注入页面，实现元素选中、高亮、属性编辑等交互能力；源代码变更后 HMR 热更新自动刷新页面", 11, RGBColor(80, 85, 105))

ARROW_COLOR = RGBColor(180, 185, 200)
T = Inches(0.03)
H = Inches(0.13)
gap_top = dia_y + box_h
gap_bot = page_y

# ── Vertical arrow: Plugin ↓ Page (Content Script) ──
ax1 = plug_bx + box_w / 2
sv1 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE,
      ax1 - T / 2, gap_top, T, gap_bot - gap_top - H)
sv1.fill.solid(); sv1.fill.fore_color.rgb = ARROW_COLOR; sv1.line.fill.background()
tri1 = s3.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
       ax1 - H / 2, gap_bot - H, H, H)
tri1.fill.solid(); tri1.fill.fore_color.rgb = ARROW_COLOR
tri1.line.fill.background(); tri1.rotation = 180.0
text_box(s3, ax1 + Inches(0.1), (gap_top + gap_bot) / 2 - Inches(0.12),
         Inches(1.2), Inches(0.25), "Content Script", 9, GRAY)

# ── Vertical arrow: Code ↓ Page (HMR) ──
ax2 = code_bx + box_w / 2
sv2 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE,
      ax2 - T / 2, gap_top, T, gap_bot - gap_top - H)
sv2.fill.solid(); sv2.fill.fore_color.rgb = ARROW_COLOR; sv2.line.fill.background()
tri2 = s3.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
       ax2 - H / 2, gap_bot - H, H, H)
tri2.fill.solid(); tri2.fill.fore_color.rgb = ARROW_COLOR
tri2.line.fill.background(); tri2.rotation = 180.0
text_box(s3, ax2 + Inches(0.1), (gap_top + gap_bot) / 2 - Inches(0.12),
         Inches(1.8), Inches(0.25), "文件变更 → HMR", 9, GRAY)


# ── Save ──
out = "/Users/vengeanliu/Documents/ai-workspace/PrismDesign/docs/PrismDesign-介绍.pptx"
prs.save(out)
print(f"Saved: {out}")
